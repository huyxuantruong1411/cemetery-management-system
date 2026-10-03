import io
from decimal import Decimal

import openpyxl
from fastapi.testclient import TestClient

from app.db.session import SessionLocal
from app.main import app
from app.modules.auth.models import User
from app.modules.contracts.models import Contract
from app.modules.documents.models import DocumentVersion
from app.modules.profiles.models import Customer
from app.services.excel_service import ExcelService

client = TestClient(app)

# Helper to get admin token
def get_admin_token() -> str:
    resp = client.post(
        "/api/v1/auth/login",
        json={"username": "admin", "password": "Admin2026!"},
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["access_token"]


def test_upload_valid_pdf_and_download():
    """Verify uploading a valid PDF document with magic bytes and downloading it with auth."""
    token = get_admin_token()
    valid_pdf_content = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF"

    # Upload valid PDF
    upload_resp = client.post(
        "/api/v1/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test_contract.pdf", valid_pdf_content, "application/pdf")},
    )
    assert upload_resp.status_code == 201, upload_resp.text
    file_data = upload_resp.json()
    assert file_data["file_name"] == "test_contract.pdf"
    assert file_data["mime_type"] == "application/pdf"
    assert file_data["state"] == "READY"
    file_id = file_data["file_id"]

    # Download document with auth
    download_resp = client.get(
        f"/api/v1/documents/{file_id}/download",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert download_resp.status_code == 200
    assert download_resp.content == valid_pdf_content
    assert download_resp.headers["content-type"] == "application/pdf"
    assert "X-SHA256" in download_resp.headers

    # Preview document
    preview_resp = client.get(
        f"/api/v1/documents/{file_id}/preview",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert preview_resp.status_code == 200
    assert preview_resp.content == valid_pdf_content


def test_upload_spoofed_file_rejected():
    """Verify that uploading a text file renamed to .pdf without %PDF- magic bytes is rejected."""
    token = get_admin_token()
    fake_pdf_content = b"This is plain text pretending to be a PDF file."

    upload_resp = client.post(
        "/api/v1/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("malicious.pdf", fake_pdf_content, "application/pdf")},
    )
    assert upload_resp.status_code == 400
    assert "Định dạng tệp không được hỗ trợ" in upload_resp.json()["detail"]


def test_download_unauthorized_rejected():
    """Verify that unauthenticated download requests are rejected with 401."""
    resp = client.get("/api/v1/documents/non_existent_file/download")
    assert resp.status_code == 401


def test_document_versioning():
    """Verify linking files to a contract and automatic version numbering."""
    token = get_admin_token()
    db = SessionLocal()

    customer = None
    contract = None
    try:
        # Create test customer & contract
        customer = Customer(
            customer_code="C_DOC_TEST",
            full_name="Khách Hàng Chứng Từ",
            citizen_id="079099887766",
            phone_number="0909999888",
            address="Hà Nội",
        )
        db.add(customer)
        db.flush()

        admin_user = db.query(User).filter(User.username == "admin").first()
        admin_id = admin_user.user_id if admin_user else None

        contract = Contract(
            contract_code="HD_DOC_001",
            contract_type="LAND_PURCHASE",
            customer_id=customer.customer_id,
            status="DRAFT",
            total_amount=Decimal("50000000.00"),
            created_by_user_id=admin_id,
        )
        db.add(contract)
        db.commit()

        # Upload 2 versions
        pdf_v1 = b"%PDF-1.4\nVersion 1 content\n%%EOF"
        pdf_v2 = b"%PDF-1.4\nVersion 2 updated content\n%%EOF"

        up1 = client.post(
            "/api/v1/documents/upload",
            headers={"Authorization": f"Bearer {token}"},
            files={"file": ("contract_v1.pdf", pdf_v1, "application/pdf")},
        )
        file_id_1 = up1.json()["file_id"]

        up2 = client.post(
            "/api/v1/documents/upload",
            headers={"Authorization": f"Bearer {token}"},
            files={"file": ("contract_v2.pdf", pdf_v2, "application/pdf")},
        )
        file_id_2 = up2.json()["file_id"]

        # Link Version 1
        link1 = client.post(
            "/api/v1/documents/versions",
            headers={"Authorization": f"Bearer {token}"},
            json={
                "file_id": file_id_1,
                "document_type": "SIGNED_CONTRACT",
                "contract_id": contract.contract_id,
                "notes": "Bản quét lần 1",
            },
        )
        assert link1.status_code == 201
        assert link1.json()["version_no"] == 1

        # Link Version 2
        link2 = client.post(
            "/api/v1/documents/versions",
            headers={"Authorization": f"Bearer {token}"},
            json={
                "file_id": file_id_2,
                "document_type": "SIGNED_CONTRACT",
                "contract_id": contract.contract_id,
                "notes": "Bản quét ký lại lần 2",
            },
        )
        assert link2.status_code == 201
        assert link2.json()["version_no"] == 2

    finally:
        db.rollback()
        if contract and contract.contract_id:
            db.query(DocumentVersion).filter(DocumentVersion.contract_id == contract.contract_id).delete()
            db.query(Contract).filter(Contract.contract_id == contract.contract_id).delete()
        if customer and customer.customer_id:
            db.query(Customer).filter(Customer.customer_id == customer.customer_id).delete()
        db.commit()
        db.close()


def test_generate_sample_vietnamese_pdf():
    """Verify Vietnamese PDF generation produces a valid PDF file in MinIO."""
    token = get_admin_token()

    resp = client.post(
        "/api/v1/documents/sample-contract-pdf",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert resp.status_code == 201, resp.text
    data = resp.json()
    assert data["mime_type"] == "application/pdf"
    assert data["file_size_bytes"] > 1000  # Valid non-empty PDF

    # Download and verify PDF magic bytes
    dl = client.get(
        f"/api/v1/documents/{data['file_id']}/download",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert dl.status_code == 200
    assert dl.content.startswith(b"%PDF-")


def test_excel_service_generation():
    """Verify Excel report generation produces a readable, styled workbook."""
    headers = ["Mã Ô Mộ", "Khu Vực", "Chủ Sở Hữu", "Trạng Thái", "Đơn Giá (VNĐ)"]
    rows = [
        ["A1-01", "Khu Vạn Phúc", "Nguyễn Văn An", "OCCUPIED", 120000000],
        ["A1-02", "Khu Vạn Phúc", "Trần Thị Bình", "RESERVED", 120000000],
        ["A1-03", "Khu Vạn Phúc", "Chưa Bán", "EMPTY_UNSOLD", 120000000],
    ]

    xlsx_bytes = ExcelService.generate_report_xlsx(
        sheet_title="Báo Cáo Ô Mộ",
        headers=headers,
        rows=rows,
    )
    assert len(xlsx_bytes) > 0

    # Parse with openpyxl to verify correctness
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws.title == "Báo Cáo Ô Mộ"
    assert ws.cell(row=1, column=1).value == "Mã Ô Mộ"
    assert ws.cell(row=2, column=1).value == "A1-01"
    assert ws.cell(row=2, column=5).value == 120000000
