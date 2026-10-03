import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def sanitize_excel_cell(val: Any) -> Any:
    """Prevent CSV/Excel Formula Injection by prepending single quote to formula trigger characters."""
    if isinstance(val, str) and val and val[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{val}"
    return val


class ExcelService:
    sanitize_excel_cell = staticmethod(sanitize_excel_cell)

    @classmethod
    def generate_report_xlsx(
        cls,
        sheet_title: str,
        headers: list[str],
        rows: list[list[Any]],
    ) -> bytes:
        """Generate an Excel workbook with styled headers, data rows, and formula injection protection."""
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title[:31]

        # Styling
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="24594D", end_color="24594D", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Data rows
        data_font = Font(name="Arial", size=10)
        for row_data in rows:
            sanitized_row = [cls.sanitize_excel_cell(c) for c in row_data]
            ws.append(sanitized_row)
            current_row = ws.max_row
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_idx)
                cell.font = data_font

        # Auto-adjust column width
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
